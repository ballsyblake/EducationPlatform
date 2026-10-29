"""Write the coach assessment outcome sheets as a Word document.

    python3 scripts/export-assessment-docx.py prisma/data/b-diploma-2026.json \
        --course b-diploma-moreton-city-excelsior-2026 -o sheets.docx

Reads the JSON the register extractors produce — `extract-b-diploma-2026.py`
and its C Diploma sibling — and lays each coach out on a page of their own, the
way FQ's own `C - LICENCE COACH ASSESSMENT` file does. Their copy was generated
by python-docx too, which is why this reads as a sibling of their document
rather than a conversion of ours.

**This is a second rendering of a layout that already exists**, in
`src/components/outcome-sheet.tsx`, and the two will drift. That is a cost
taken on purpose: the app's sheet is the one people read on screen and print,
and this one exists because FQ circulates .docx files that other people edit
and sign. Keep this script for the Word file and change the component for
everything else; when a field moves, move it in both.

What it cannot print is any outcome. The registers hold no six-area scores —
the Assessment tab ends in a free-text `Session Rating: n` on the old 1-5 scale
— so the Highly Competent / Competent / Not Yet Competent boxes come out empty
and the sheet says why. That is the register's gap, not this script's, and
filling it in here would be inventing a conversion nobody has published.

**The registers are real people.** Names are real and the extractors anonymise
only the addresses. Whatever this writes goes where the register went and
nowhere else.
"""
import argparse, json, re, sys

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Pt, RGBColor, Cm

# Football Queensland's Maroon, from the brand guidelines (2.B) and matching
# `--color-maroon-800` in globals.css, so the Word file and the page are the
# same document in two formats rather than two documents.
MAROON = RGBColor(0x62, 0x13, 0x33)
INK = RGBColor(0x3D, 0x3D, 0x3D)
MUTED = RGBColor(0x6B, 0x6B, 0x6B)

# Football Australia's rubric, transcribed from `src/lib/delivery-rubric.ts`.
# Kept in step by hand, which is the price of the second rendering above.
AREAS = ["Objective", "Content", "Organisation", "Presenting", "Coaching", "Environment"]
LEVELS = [
    (0, "Not Evident", "The required elements are absent, vague, or entirely ineffective."),
    (1, "Developing", "Elements are present but lack consistency, clarity, or full alignment with objectives."),
    (2, "Proficient", "Elements are clear, consistent, and effectively support player understanding and engagement."),
    (3, "Expertise", "Elements are exemplary, seamlessly integrated, and maximise player learning and engagement."),
]
MAX_POINTS = len(AREAS) * 3

# `DELIVERY_THRESHOLDS`. The B and A Highly Competent lines are inferred from
# the C Diploma's three-point step, not quoted from FA — see the rubric module.
THRESHOLDS = {"C": (9, 12, False), "B": (12, 15, True), "A": (15, 18, True)}
LEVEL_PATTERN = re.compile(r"\b([abc])\s*[-/]?\s*(?:diploma|licence|license|certificate)\b", re.I)

OUTCOMES = ["Highly Competent", "Competent", "Not Yet Competent"]


def diploma_level(*sources):
    for s in sources:
        m = LEVEL_PATTERN.search(s or "")
        if m:
            return m.group(1).upper()
    return None


def plan_steps(action_plan):
    """The numbering the assessor wrote, taken back off so the list owns it."""
    if not action_plan:
        return []
    steps = []
    for line in action_plan.splitlines():
        line = re.sub(r"^\s*\d+[.)]\s*", "", line).strip()
        if line:
            steps.append(line)
    # Some assessors number with a bare digit and no stop — "1 Be clear with
    # the game rules". Taken off only where the digit is the step's own
    # position, so an action that opens with a number keeps it.
    return [re.sub(rf"^{i + 1}\s+(?=\D)", "", s) for i, s in enumerate(steps)]


# --------------------------------------------------------------------------
# Small helpers over python-docx, which has no notion of a styled run of text
# --------------------------------------------------------------------------


def run(paragraph, text, *, bold=False, size=9, colour=INK, caps=False):
    r = paragraph.add_run(text.upper() if caps else text)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.color.rgb = colour
    return r


def para(container, text="", **kw):
    p = container.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    if text:
        run(p, text, **kw)
    return p


def heading(doc, text):
    """The maroon section band FQ's form rules each block off with."""
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    cell = t.rows[0].cells[0]
    cell.paragraphs[0].paragraph_format.space_after = Pt(0)
    run(cell.paragraphs[0], text, bold=True, size=8, colour=MAROON, caps=True)
    return t


def facts_table(doc, delivery):
    """The form's 2x3 block: assessor, block, component, topic, date, session."""
    fields = [
        ("Assessor", delivery.get("assessor")),
        ("Block / when", delivery.get("block")),
        ("Component", delivery.get("component")),
        ("Topic", delivery.get("topic")),
        # Neither is in the registers: their "When" is the block and nothing
        # dates the session. They print as the form's own blank.
        ("Date", delivery.get("date")),
        ("Session", delivery.get("session")),
    ]
    t = doc.add_table(rows=3, cols=2)
    t.style = "Table Grid"
    for i, (label, value) in enumerate(fields):
        cell = t.rows[i // 2].cells[i % 2]
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run(p, f"{label}: ", bold=True, size=8, colour=MUTED, caps=True)
        run(p, value or "", size=9)
    return t


def scoresheet(doc, scores):
    """The six areas and their marks. Only reached where a register has them."""
    t = doc.add_table(rows=len(AREAS) + 1, cols=2)
    t.style = "Table Grid"
    total = 0
    for i, area in enumerate(AREAS):
        score = scores.get(area)
        if isinstance(score, int):
            total += score
        cells = t.rows[i].cells
        para(cells[0], area, size=9)
        # A blank area is not a nought — the total is incomplete until all six
        # carry a mark, same rule as `totalPoints`.
        para(cells[1], f"{score if isinstance(score, int) else '—'} / 3", size=9)
    cells = t.rows[-1].cells
    para(cells[0], "Total", bold=True, size=8, caps=True)
    para(cells[1], f"{total} / {MAX_POINTS}", bold=True, size=9)
    return total


def boxes(doc, ticked=None):
    """The outcome row. An empty box is a box somebody still has to tick."""
    p = para(doc)
    p.paragraph_format.space_before = Pt(4)
    for label in OUTCOMES:
        run(p, "[X] " if label == ticked else "[  ] ", bold=True, size=10, colour=MAROON)
        run(p, f"{label}    ", bold=label == ticked, size=9)


def sheet(doc, coach, course, first):
    """One coach, one page."""
    if not first:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    qualification = course.get("qualification") or course.get("title") or ""
    level = diploma_level(qualification, course.get("title"))
    competent, highly, inferred = THRESHOLDS.get(level or "C", THRESHOLDS["C"])

    header = para(doc)
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run(header, f"{qualification} coach assessment", bold=True, size=10, colour=MAROON, caps=True)

    name = " ".join(x for x in [coach.get("firstName"), coach.get("lastName")] if x)
    para(doc, name, bold=True, size=16, colour=RGBColor(0, 0, 0), caps=True)
    line = para(doc)
    run(line, "Coach", size=8, colour=MUTED, caps=True)
    detail = " · ".join(x for x in [course.get("title"), coach.get("club"), course.get("location")] if x)
    run(line, f"    {detail}", size=8, colour=MUTED)

    deliveries = coach.get("deliveries") or []
    # A coach nobody has written up still gets a form, because whoever prints
    # one is taking it to the grass to fill in.
    for delivery in deliveries or [{}]:
        if len(deliveries) > 1:
            p = para(doc, f"Delivery {delivery.get('deliveryNo')}", bold=True, size=8, colour=MAROON, caps=True)
            p.paragraph_format.space_before = Pt(8)
        facts_table(doc, delivery)

        scores = delivery.get("scores")
        if scores:
            heading(doc, f"Assessment — six areas, {MAX_POINTS} points")
            scoresheet(doc, scores)

        if delivery.get("rating") is not None:
            note = f"Rated {float(delivery['rating']):.1f} / 5 on the register."
            if not scores:
                note += " This session has not been scored on the six areas."
            para(doc, note, size=8, colour=MUTED)

        heading(doc, "Assessor comment / evidence")
        body = doc.add_table(rows=1, cols=1)
        body.style = "Table Grid"
        cell = body.rows[0].cells[0]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        run(cell.paragraphs[0], delivery.get("comment") or "", size=9)

        heading(doc, "Action plan")
        steps = plan_steps(delivery.get("actionPlan"))
        # Three rows whether or not three were written: a plan with room left
        # on it reads as a plan rather than a finished list.
        plan = doc.add_table(rows=3, cols=1)
        plan.style = "Table Grid"
        for i in range(3):
            p = plan.rows[i].cells[0].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run(p, f"{i + 1}. ", bold=True, size=9, colour=MUTED)
            run(p, steps[i] if i < len(steps) else "", size=9)

    # The thresholds stay on every sheet — they are what makes the outcome
    # self-evidencing — but the four-rung table does not. It is identical on
    # all thirty-five pages and tall enough to tip each one onto a second,
    # which is thirty-five pages carrying a rating guide somebody has already
    # read. It opens the document instead; see `cover`.
    scale = (
        f"Six areas, {MAX_POINTS} points in all, scored on a best fit of what the session showed "
        f"and never averaged. On a {qualification}: under {competent} is Not Yet Competent, "
        f"{competent} to {highly - 1} is Competent, and {highly} or more is Highly Competent."
    )
    if inferred:
        scale += (
            " Football Australia has published the competency line for this diploma but not the "
            "Highly Competent one, which follows the C Diploma's three-point step."
        )
    para(doc, scale, size=8, colour=MUTED)

    scored = [d for d in deliveries if d.get("scores")]
    heading(doc, "Overall")
    if scored:
        total = sum(v for v in scored[-1]["scores"].values() if isinstance(v, int))
        complete = len([v for v in scored[-1]["scores"].values() if isinstance(v, int)]) == len(AREAS)
        label = None
        if complete:
            label = "Highly Competent" if total >= highly else "Competent" if total >= competent else "Not Yet Competent"
        para(doc, f"{total} / {MAX_POINTS}", bold=True, size=14, colour=MAROON)
        boxes(doc, label)
        if not complete:
            # A total off a half-filled sheet is not an outcome, and a reader
            # given the number with no box ticked would otherwise be left
            # guessing whether the blanks were noughts.
            para(
                doc,
                "Not every area on this session has been scored, so no outcome can be read. "
                "A blank area is not a nought.",
                size=8,
                colour=MUTED,
            )
        elif len(scored) > 1:
            para(
                doc,
                f"Read off delivery {scored[-1].get('deliveryNo')}, the most recent session "
                "scored in full. Totals are never averaged across sessions.",
                size=8,
                colour=MUTED,
            )
    else:
        para(doc, f"— / {MAX_POINTS}", bold=True, size=14, colour=MAROON)
        boxes(doc)
        para(
            doc,
            "This register holds no six-area scores, so there is no outcome to read. The boxes "
            "above are for an assessor to complete.",
            size=8,
            colour=MUTED,
        )

    if coach.get("rating") is not None:
        para(
            doc,
            f"Separately, the course rating on the register is {float(coach['rating']):.1f} / 5. "
            "It is a judgement across everything this coach delivered rather than a total of "
            "their sessions, and it does not convert to the points above.",
            size=8,
            colour=MUTED,
        )

    # One paragraph rather than a table. A two-cell table with room to sign in
    # it was tall enough to tip every sheet onto a second page, which on a
    # thirty-five coach cohort is thirty-five pages carrying a signature rule
    # and nothing else.
    sign = para(doc)
    sign.paragraph_format.space_before = Pt(10)
    run(sign, "______________________________  ", size=9, colour=MUTED)
    run(sign, "Assessor signature and date", size=8, colour=MUTED)
    run(sign, "      ______________________________  ", size=9, colour=MUTED)
    run(sign, "Coach acknowledgement and date", size=8, colour=MUTED)


def cover(doc, course, competent, highly, inferred):
    """The front page: what this file is, and the rubric it is scored on."""
    qualification = course.get("qualification") or course.get("title") or ""
    para(doc, f"{qualification} coach assessment", bold=True, size=18, colour=MAROON, caps=True)
    para(doc, course.get("title") or "", size=10)
    detail = " · ".join(x for x in [course.get("address") or course.get("location"), course.get("season")] if x)
    para(doc, detail, size=9, colour=MUTED)

    staff = ", ".join(s["name"] for s in course.get("staff") or [] if s.get("name"))
    if staff:
        p = para(doc)
        p.paragraph_format.space_before = Pt(8)
        run(p, "Course team: ", bold=True, size=9, colour=MUTED)
        run(p, staff, size=9)

    p = para(doc)
    p.paragraph_format.space_before = Pt(16)
    run(p, "Assessment rating guide", bold=True, size=10, colour=MAROON, caps=True)

    guide = doc.add_table(rows=len(LEVELS) + 1, cols=3)
    guide.style = "Table Grid"
    for i, text in enumerate(["Score", "Level", "What it means"]):
        para(guide.rows[0].cells[i], text, bold=True, size=9)
    for i, (score, label, description) in enumerate(LEVELS, start=1):
        para(guide.rows[i].cells[0], str(score), bold=True, size=9)
        para(guide.rows[i].cells[1], label, size=9)
        para(guide.rows[i].cells[2], description, size=9)

    note = (
        f"Six areas — {', '.join(AREAS)} — scored 0 to 3 on a best fit of what the session "
        f"showed, {MAX_POINTS} points in all and never averaged. A blank area is not a nought: "
        f"the total is incomplete until all six carry a mark. On a {qualification}: under "
        f"{competent} is Not Yet Competent, {competent} to {highly - 1} is Competent, and "
        f"{highly} or more is Highly Competent."
    )
    if inferred:
        note += (
            " Football Australia has published the competency line for this diploma but not the "
            "Highly Competent one, which follows the C Diploma's three-point step."
        )
    p = para(doc, note, size=9, colour=MUTED)
    p.paragraph_format.space_before = Pt(8)


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("data", help="JSON from one of the register extractors")
    ap.add_argument("--course", help="course slug; omit to take the only one, or list them")
    ap.add_argument("-o", "--out", required=True)
    args = ap.parse_args()

    courses = json.load(open(args.data, encoding="utf-8"))["courses"]
    if args.course:
        courses = [c for c in courses if c["slug"] == args.course]
    if len(courses) != 1:
        print("Name one course with --course:", file=sys.stderr)
        for c in json.load(open(args.data, encoding="utf-8"))["courses"]:
            print(f"  {c['slug']}", file=sys.stderr)
        return 1
    course = courses[0]

    doc = Document()
    section = doc.sections[0]
    for attr in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, attr, Cm(1.4))

    # Gibson is the brand typeface and is licensed, so it is not here; the
    # guidelines name Arial as the substitute, which is what the app's own
    # `--font-brand` falls back to. Set on Normal so every table cell inherits
    # it — python-docx otherwise leaves them on the template's serif.
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)

    # The register's own order: the main roster, then anybody catching up.
    roster = [c for c in course["coaches"] if c.get("track") != "CATCH_UP"]
    roster += [c for c in course["coaches"] if c.get("track") == "CATCH_UP"]
    level = diploma_level(course.get("qualification"), course.get("title"))
    cover(doc, course, *THRESHOLDS.get(level or "C", THRESHOLDS["C"]))
    for coach in roster:
        sheet(doc, coach, course, first=False)

    doc.save(args.out)
    scored = sum(1 for c in roster for d in (c.get("deliveries") or []) if d.get("scores"))
    print(
        f"{course['title']}: {len(roster)} sheets, "
        f"{sum(len(c.get('deliveries') or []) for c in roster)} deliveries, "
        f"{scored} scored on the rubric"
    )
    print(f"wrote {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
